from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path
from typing import Any
import re
import sqlite3

from fastapi import APIRouter, HTTPException
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field


# ============================================================
# OPTIONAL EXPORT DEPENDENCIES
# ============================================================

try:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer,
        Table,
        TableStyle,
        PageBreak,
        KeepTogether,
    )

    REPORTLAB_AVAILABLE = True

except Exception:
    REPORTLAB_AVAILABLE = False


try:
    from openpyxl import Workbook
    from openpyxl.styles import (
        Alignment,
        Border,
        Font,
        PatternFill,
        Side,
    )
    from openpyxl.utils import get_column_letter

    OPENPYXL_AVAILABLE = True

except Exception:
    OPENPYXL_AVAILABLE = False


# ============================================================
# ROUTER
# ============================================================

router = APIRouter(
    prefix="/api/report-pack",
    tags=["Monthly Report Pack"],
)


# ============================================================
# PATHS
# ============================================================

PROJECT_ROOT = (
    Path(__file__)
    .resolve()
    .parents[2]
)

DATABASE_PATH = (
    PROJECT_ROOT
    / "data"
    / "database"
    / "controlsight.db"
)

REPORT_OUTPUT_DIR = (
    PROJECT_ROOT
    / "data"
    / "report_packs"
)

REPORT_OUTPUT_DIR.mkdir(
    parents=True,
    exist_ok=True,
)


# ============================================================
# CONSTANTS
# ============================================================

PERIOD_PATTERN = re.compile(
    r"^\d{4}-(0[1-9]|1[0-2])$"
)

SUPPORTED_FORMATS = {
    "pdf",
    "xlsx",
}


# ============================================================
# REQUEST MODEL
# ============================================================

class ReportPackRequest(BaseModel):

    project_code: str = Field(
        min_length=1,
        max_length=100,
    )

    reporting_period: str | None = None

    formats: list[str] = Field(
        default_factory=lambda: [
            "pdf",
            "xlsx",
        ]
    )


# ============================================================
# DATABASE
# ============================================================

def get_connection():

    if not DATABASE_PATH.exists():

        raise HTTPException(
            status_code=500,
            detail=(
                "ControlSight database not found: "
                f"{DATABASE_PATH}"
            ),
        )

    connection = sqlite3.connect(
        DATABASE_PATH
    )

    connection.row_factory = sqlite3.Row

    return connection


def table_exists(
    connection,
    table_name: str,
):

    row = connection.execute(
        """
        SELECT name

        FROM sqlite_master

        WHERE
            type IN ('table', 'view')
            AND
            name = ?
        """,
        (
            table_name,
        ),
    ).fetchone()

    return bool(row)


def table_columns(
    connection,
    table_name: str,
):

    if not table_exists(
        connection,
        table_name,
    ):
        return set()

    rows = connection.execute(
        f"""
        PRAGMA table_info(
            {table_name}
        )
        """
    ).fetchall()

    return {
        row["name"]
        for row in rows
    }


def first_existing_column(
    connection,
    table_name: str,
    candidates: list[str],
):

    columns = table_columns(
        connection,
        table_name,
    )

    for candidate in candidates:

        if candidate in columns:
            return candidate

    return None


# ============================================================
# GENERIC HELPERS
# ============================================================

def safe_float(
    value,
    default=0.0,
):

    try:
        if value is None:
            return default

        return float(value)

    except Exception:
        return default


def safe_int(
    value,
    default=0,
):

    try:
        if value is None:
            return default

        return int(value)

    except Exception:
        return default


def money(
    value,
):

    number = safe_float(
        value
    )

    if number < 0:

        return (
            f"-£{abs(number):,.0f}"
        )

    return (
        f"£{number:,.0f}"
    )


def percentage(
    value,
):

    return (
        f"{safe_float(value):,.2f}%"
    )


def text_value(
    value,
):

    if value is None:
        return "—"

    text = str(value).strip()

    return text if text else "—"


def safe_filename_part(
    value: str,
):

    cleaned = re.sub(
        r"[^A-Za-z0-9_-]+",
        "_",
        str(value),
    )

    return cleaned.strip(
        "_"
    )


def utc_timestamp():

    return (
        datetime.now(UTC)
        .replace(
            microsecond=0
        )
        .isoformat()
        +
        "Z"
    )


# ============================================================
# PROJECT
# ============================================================

def get_project_scorecard(
    connection,
    project_code,
    reporting_period=None,
):

    # --------------------------------------------------------
    # BASE PROJECT / LATEST MANAGEMENT SCORECARD
    # --------------------------------------------------------

    row = connection.execute(
        """
        SELECT *

        FROM report_management_scorecard

        WHERE project_code = ?
        """,
        (
            project_code,
        ),
    ).fetchone()


    if row is None:

        raise HTTPException(
            status_code=404,
            detail=(
                f"Project {project_code} "
                "not found."
            ),
        )


    scorecard = dict(
        row
    )


    if not reporting_period:

        return scorecard


    # --------------------------------------------------------
    # PERIOD COST SNAPSHOT
    #
    # Mirrors the historical logic used by main.py:
    #
    # 1. aggregate cost by reporting date
    # 2. calculate cumulative actual / committed
    # 3. return the final snapshot inside selected YYYY-MM
    # --------------------------------------------------------

    cost = connection.execute(
        """
        WITH monthly_cost AS (

            SELECT
                fc.project_code,
                dd.date,
                dd.year_month,

                SUM(
                    fc.planned_cost
                ) AS monthly_planned_cost,

                SUM(
                    fc.actual_cost
                ) AS monthly_actual_cost,

                SUM(
                    fc.committed_cost
                ) AS monthly_committed_cost,

                SUM(
                    fc.forecast_at_completion
                ) AS forecast_at_completion

            FROM fact_cost fc

            JOIN dim_date dd
                ON fc.date_key
                = dd.date_key

            WHERE
                fc.project_code = ?
                AND
                dd.year_month <= ?

            GROUP BY
                fc.project_code,
                dd.date,
                dd.year_month
        ),

        cumulative AS (

            SELECT
                project_code,
                date,
                year_month,

                SUM(
                    monthly_planned_cost
                ) OVER (
                    ORDER BY date
                )
                    AS cumulative_planned_cost,

                SUM(
                    monthly_actual_cost
                ) OVER (
                    ORDER BY date
                )
                    AS cumulative_actual_cost,

                SUM(
                    monthly_committed_cost
                ) OVER (
                    ORDER BY date
                )
                    AS cumulative_committed_cost,

                forecast_at_completion

            FROM monthly_cost
        )

        SELECT *

        FROM cumulative

        WHERE year_month = ?

        ORDER BY date DESC

        LIMIT 1
        """,
        (
            project_code,
            reporting_period,
            reporting_period,
        ),
    ).fetchone()


    if cost is None:

        raise HTTPException(
            status_code=404,
            detail=(
                "No cost data for "
                f"{project_code} "
                "in "
                f"{reporting_period}."
            ),
        )


    cost = dict(
        cost
    )


    scorecard[
        "actual_cost_to_date"
    ] = round(
        safe_float(
            cost.get(
                "cumulative_actual_cost"
            )
        ),
        2,
    )


    scorecard[
        "committed_cost_to_date"
    ] = round(
        safe_float(
            cost.get(
                "cumulative_committed_cost"
            )
        ),
        2,
    )


    scorecard[
        "forecast_at_completion"
    ] = round(
        safe_float(
            cost.get(
                "forecast_at_completion"
            )
        ),
        2,
    )


    # --------------------------------------------------------
    # RECOMPUTE FORECAST VARIANCE FOR SELECTED PERIOD
    # --------------------------------------------------------

    budget = safe_float(
        scorecard.get(
            "approved_budget"
        )
    )


    fac = safe_float(
        scorecard.get(
            "forecast_at_completion"
        )
    )


    forecast_variance = (
        budget
        -
        fac
    )


    scorecard[
        "forecast_variance"
    ] = round(
        forecast_variance,
        2,
    )


    scorecard[
        "forecast_variance_pct"
    ] = round(
        (
            forecast_variance
            /
            budget
            *
            100
        )
        if budget
        else 0.0,
        2,
    )


    # --------------------------------------------------------
    # PERIOD PROGRESS SNAPSHOT
    #
    # Progress is a point-in-time percentage, so unlike cost
    # it must NOT be cumulatively summed.
    # --------------------------------------------------------

    progress = connection.execute(
        """
        SELECT
            fp.project_code,
            dd.date,
            dd.year_month,
            fp.planned_progress_pct,
            fp.actual_progress_pct,
            fp.progress_variance_pct,
            fp.behind_schedule_flag

        FROM fact_progress fp

        JOIN dim_date dd
            ON fp.date_key
            = dd.date_key

        WHERE
            fp.project_code = ?
            AND
            dd.year_month = ?

        ORDER BY
            dd.date DESC

        LIMIT 1
        """,
        (
            project_code,
            reporting_period,
        ),
    ).fetchone()


    if progress is not None:

        progress = dict(
            progress
        )


        scorecard[
            "planned_progress_pct"
        ] = safe_float(
            progress.get(
                "planned_progress_pct"
            )
        )


        scorecard[
            "actual_progress_pct"
        ] = safe_float(
            progress.get(
                "actual_progress_pct"
            )
        )


        scorecard[
            "progress_variance_pct"
        ] = safe_float(
            progress.get(
                "progress_variance_pct"
            )
        )


    # --------------------------------------------------------
    # PERIOD FORECAST CONCERN CROSS-CHECK
    # --------------------------------------------------------

    actual_cost = safe_float(
        scorecard.get(
            "actual_cost_to_date"
        )
    )


    actual_progress = safe_float(
        scorecard.get(
            "actual_progress_pct"
        )
    )


    fac = safe_float(
        scorecard.get(
            "forecast_at_completion"
        )
    )


    progress_implied_eac = 0.0


    if actual_progress > 0:

        progress_implied_eac = (
            actual_cost
            /
            (
                actual_progress
                /
                100
            )
        )


    credibility_gap = (
        progress_implied_eac
        -
        fac
    )


    credibility_gap_pct = (
        (
            credibility_gap
            /
            fac
            *
            100
        )
        if fac
        else 0.0
    )


    abs_gap = abs(
        credibility_gap_pct
    )


    if abs_gap >= 20:

        credibility_status = (
            "Critical"
        )

    elif abs_gap >= 10:

        credibility_status = (
            "High"
        )

    elif abs_gap >= 5:

        credibility_status = (
            "Watch"
        )

    else:

        credibility_status = (
            "Stable"
        )


    scorecard[
        "progress_implied_eac"
    ] = round(
        progress_implied_eac,
        2,
    )


    scorecard[
        "forecast_credibility_gap"
    ] = round(
        credibility_gap,
        2,
    )


    scorecard[
        "forecast_credibility_gap_pct"
    ] = round(
        credibility_gap_pct,
        2,
    )


    scorecard[
        "forecast_credibility_status"
    ] = credibility_status


    return scorecard




# ============================================================
# CHANGE CONTROL
# ============================================================

def get_change_control(
    connection,
    project_code,
    scorecard,
):

    source_budget = 0.0

    for column in [
        "budget",
        "approved_budget",
        "current_budget",
        "total_budget",
        "project_budget",
    ]:

        if column in scorecard:

            source_budget = safe_float(
                scorecard.get(
                    column
                )
            )

            break


    original_budget = (
        source_budget
    )


    if table_exists(
        connection,
        "project_budget_baseline",
    ):

        row = connection.execute(
            """
            SELECT *

            FROM project_budget_baseline

            WHERE project_code = ?
            """,
            (
                project_code,
            ),
        ).fetchone()

        if row:

            original_budget = (
                safe_float(
                    row[
                        "original_budget"
                    ]
                )
            )


    changes = []


    if table_exists(
        connection,
        "project_change_log",
    ):

        rows = connection.execute(
            """
            SELECT *

            FROM project_change_log

            WHERE project_code = ?

            ORDER BY created_at DESC
            """,
            (
                project_code,
            ),
        ).fetchall()

        changes = [
            dict(row)
            for row in rows
        ]


    approved = sum(
        safe_float(
            row.get(
                "change_amount"
            )
        )
        for row in changes
        if row.get(
            "status"
        )
        ==
        "Approved"
    )


    pending = sum(
        safe_float(
            row.get(
                "change_amount"
            )
        )
        for row in changes
        if row.get(
            "status"
        )
        ==
        "Pending"
    )


    current_budget = (
        original_budget
        +
        approved
    )


    return {

        "original_budget":
            round(
                original_budget,
                2,
            ),

        "approved_changes":
            round(
                approved,
                2,
            ),

        "current_budget":
            round(
                current_budget,
                2,
            ),

        "pending_changes":
            round(
                pending,
                2,
            ),

        "potential_budget":
            round(
                current_budget
                +
                pending,
                2,
            ),

        "changes":
            changes,
    }


# ============================================================
# CONTINGENCY / RISK
# ============================================================

def get_risk_control(
    connection,
    project_code,
):

    contingency = 0.0


    if table_exists(
        connection,
        "project_contingency_baseline",
    ):

        row = connection.execute(
            """
            SELECT *

            FROM project_contingency_baseline

            WHERE project_code = ?
            """,
            (
                project_code,
            ),
        ).fetchone()

        if row:

            contingency = safe_float(
                row[
                    "contingency_budget"
                ]
            )


    risks = []


    if table_exists(
        connection,
        "project_risk_drawdown",
    ):

        rows = connection.execute(
            """
            SELECT *

            FROM project_risk_drawdown

            WHERE project_code = ?

            ORDER BY created_at DESC
            """,
            (
                project_code,
            ),
        ).fetchall()

        risks = [
            dict(row)
            for row in rows
        ]


    proposed = sum(
        safe_float(
            row.get(
                "amount"
            )
        )
        for row in risks
        if row.get(
            "status"
        )
        ==
        "Proposed"
    )


    approved = sum(
        safe_float(
            row.get(
                "amount"
            )
        )
        for row in risks
        if row.get(
            "status"
        )
        ==
        "Approved"
    )


    drawn = sum(
        safe_float(
            row.get(
                "amount"
            )
        )
        for row in risks
        if row.get(
            "status"
        )
        ==
        "Drawn"
    )


    committed = (
        approved
        +
        drawn
    )


    remaining = (
        contingency
        -
        committed
    )


    utilisation = (
        (
            committed
            /
            contingency
            *
            100
        )
        if contingency > 0
        else 0.0
    )


    return {

        "contingency_budget":
            round(
                contingency,
                2,
            ),

        "proposed_exposure":
            round(
                proposed,
                2,
            ),

        "approved_not_drawn":
            round(
                approved,
                2,
            ),

        "drawn":
            round(
                drawn,
                2,
            ),

        "committed":
            round(
                committed,
                2,
            ),

        "remaining":
            round(
                remaining,
                2,
            ),

        "utilisation_pct":
            round(
                utilisation,
                2,
            ),

        "risks":
            risks,
    }


# ============================================================
# CASH FLOW
# ============================================================

def get_cashflow(
    connection,
    project_code,
):

    periods = []


    if table_exists(
        connection,
        "project_cashflow",
    ):

        rows = connection.execute(
            """
            SELECT *

            FROM project_cashflow

            WHERE project_code = ?

            ORDER BY reporting_period
            """,
            (
                project_code,
            ),
        ).fetchall()

        periods = [
            dict(row)
            for row in rows
        ]


    cumulative_forecast = 0.0
    cumulative_actual = 0.0

    calculated_periods = []


    for row in periods:

        forecast = safe_float(
            row.get(
                "forecast_cash"
            )
        )

        actual = safe_float(
            row.get(
                "actual_cash"
            )
        )

        committed = safe_float(
            row.get(
                "committed_cash"
            )
        )


        variance = (
            actual
            -
            forecast
        )


        variance_pct = (
            (
                variance
                /
                forecast
                *
                100
            )
            if forecast > 0
            else 0.0
        )


        cumulative_forecast += (
            forecast
        )

        cumulative_actual += (
            actual
        )


        calculated_periods.append(
            {

                **row,

                "variance":
                    round(
                        variance,
                        2,
                    ),

                "variance_pct":
                    round(
                        variance_pct,
                        2,
                    ),

                "cumulative_forecast":
                    round(
                        cumulative_forecast,
                        2,
                    ),

                "cumulative_actual":
                    round(
                        cumulative_actual,
                        2,
                    ),
            }
        )


    total_forecast = sum(
        safe_float(
            row.get(
                "forecast_cash"
            )
        )
        for row in periods
    )


    total_actual = sum(
        safe_float(
            row.get(
                "actual_cash"
            )
        )
        for row in periods
    )


    total_committed = sum(
        safe_float(
            row.get(
                "committed_cash"
            )
        )
        for row in periods
    )


    total_variance = (
        total_actual
        -
        total_forecast
    )


    return {

        "total_forecast":
            round(
                total_forecast,
                2,
            ),

        "total_actual":
            round(
                total_actual,
                2,
            ),

        "total_committed":
            round(
                total_committed,
                2,
            ),

        "total_variance":
            round(
                total_variance,
                2,
            ),

        "periods":
            calculated_periods,
    }


# ============================================================
# WORKFLOW COMMENTARY
# ============================================================

def get_workflow_commentary(
    connection,
    project_code,
    reporting_period,
):

    if not table_exists(
        connection,
        "workflow_commentary",
    ):

        return {}


    columns = table_columns(
        connection,
        "workflow_commentary",
    )


    if (
        "reporting_period"
        in columns
        and
        reporting_period
    ):

        row = connection.execute(
            """
            SELECT *

            FROM workflow_commentary

            WHERE
                project_code = ?
                AND
                reporting_period = ?

            ORDER BY rowid DESC

            LIMIT 1
            """,
            (
                project_code,
                reporting_period,
            ),
        ).fetchone()

    else:

        row = connection.execute(
            """
            SELECT *

            FROM workflow_commentary

            WHERE project_code = ?

            ORDER BY rowid DESC

            LIMIT 1
            """,
            (
                project_code,
            ),
        ).fetchone()


    return (
        dict(row)
        if row
        else {}
    )


# ============================================================
# WORKFLOW ACTIONS
# ============================================================

def get_workflow_actions(
    connection,
    project_code,
):

    if not table_exists(
        connection,
        "workflow_actions",
    ):

        return []


    rows = connection.execute(
        """
        SELECT *

        FROM workflow_actions

        WHERE project_code = ?

        ORDER BY
            CASE status
                WHEN 'Open' THEN 1
                WHEN 'In Progress' THEN 2
                WHEN 'On Hold' THEN 3
                WHEN 'Closed' THEN 4
                ELSE 5
            END,
            rowid DESC
        """,
        (
            project_code,
        ),
    ).fetchall()


    return [
        dict(row)
        for row in rows
    ]


# ============================================================
# SCHEDULE SUMMARY
# ============================================================

def get_schedule_summary(
    connection,
    project_code,
):

    if not table_exists(
        connection,
        "fact_schedule",
    ):

        return {}


    columns = table_columns(
        connection,
        "fact_schedule",
    )


    delayed_column = (
        "delayed_activity_flag"
        if "delayed_activity_flag"
        in columns
        else None
    )


    critical_column = (
        "is_critical"
        if "is_critical"
        in columns
        else None
    )


    total = connection.execute(
        """
        SELECT COUNT(*) AS count

        FROM fact_schedule

        WHERE project_code = ?
        """,
        (
            project_code,
        ),
    ).fetchone()


    delayed = 0
    critical = 0
    critical_delayed = 0


    if delayed_column:

        row = connection.execute(
            f"""
            SELECT COUNT(*) AS count

            FROM fact_schedule

            WHERE
                project_code = ?
                AND
                {delayed_column} = 1
            """,
            (
                project_code,
            ),
        ).fetchone()

        delayed = safe_int(
            row["count"]
        )


    if critical_column:

        row = connection.execute(
            f"""
            SELECT COUNT(*) AS count

            FROM fact_schedule

            WHERE
                project_code = ?
                AND
                {critical_column} = 1
            """,
            (
                project_code,
            ),
        ).fetchone()

        critical = safe_int(
            row["count"]
        )


    if (
        delayed_column
        and
        critical_column
    ):

        row = connection.execute(
            f"""
            SELECT COUNT(*) AS count

            FROM fact_schedule

            WHERE
                project_code = ?
                AND
                {critical_column} = 1
                AND
                {delayed_column} = 1
            """,
            (
                project_code,
            ),
        ).fetchone()

        critical_delayed = safe_int(
            row["count"]
        )


    return {

        "total_activities":
            safe_int(
                total["count"]
            ),

        "delayed_activities":
            delayed,

        "critical_activities":
            critical,

        "critical_delayed":
            critical_delayed,
    }


# ============================================================
# PROCUREMENT SUMMARY
# ============================================================

def get_procurement_summary(
    connection,
    project_code,
):

    if not table_exists(
        connection,
        "fact_procurement",
    ):

        return {}


    columns = table_columns(
        connection,
        "fact_procurement",
    )


    total = connection.execute(
        """
        SELECT COUNT(*) AS count

        FROM fact_procurement

        WHERE project_code = ?
        """,
        (
            project_code,
        ),
    ).fetchone()


    delayed = 0
    critical_delays = 0


    if (
        "delayed_delivery_flag"
        in columns
    ):

        row = connection.execute(
            """
            SELECT COUNT(*) AS count

            FROM fact_procurement

            WHERE
                project_code = ?
                AND
                delayed_delivery_flag = 1
            """,
            (
                project_code,
            ),
        ).fetchone()

        delayed = safe_int(
            row["count"]
        )


    if (
        "critical_delay_flag"
        in columns
    ):

        row = connection.execute(
            """
            SELECT COUNT(*) AS count

            FROM fact_procurement

            WHERE
                project_code = ?
                AND
                critical_delay_flag = 1
            """,
            (
                project_code,
            ),
        ).fetchone()

        critical_delays = safe_int(
            row["count"]
        )


    return {

        "total_items":
            safe_int(
                total["count"]
            ),

        "delayed_items":
            delayed,

        "critical_delays":
            critical_delays,
    }


# ============================================================
# REPORTING PERIOD
# ============================================================

def resolve_reporting_period(
    connection,
    requested_period,
    project_code=None,
):

    if requested_period:

        if not PERIOD_PATTERN.match(
            requested_period
        ):

            raise HTTPException(
                status_code=422,
                detail=(
                    "Reporting period must use "
                    "YYYY-MM format."
                ),
            )

        return requested_period


    # --------------------------------------------------------
    # DEFAULT PERIOD
    #
    # A Report Pack performance period must be driven by the
    # project's cost/progress reporting data, not by cash-flow.
    # Cash-flow may extend into future planning periods.
    # --------------------------------------------------------

    if project_code:

        row = connection.execute(
            """
            SELECT
                MAX(
                    dd.year_month
                ) AS reporting_period

            FROM fact_cost fc

            JOIN dim_date dd
                ON fc.date_key
                = dd.date_key

            WHERE
                fc.project_code = ?
            """,
            (
                project_code,
            ),
        ).fetchone()


        if (
            row
            and
            row[
                "reporting_period"
            ]
        ):

            return row[
                "reporting_period"
            ]


    # --------------------------------------------------------
    # GLOBAL FALLBACK
    # --------------------------------------------------------

    row = connection.execute(
        """
        SELECT
            MAX(
                dd.year_month
            ) AS reporting_period

        FROM fact_cost fc

        JOIN dim_date dd
            ON fc.date_key
            = dd.date_key
        """
    ).fetchone()


    if (
        row
        and
        row[
            "reporting_period"
        ]
    ):

        return row[
            "reporting_period"
        ]


    return (
        datetime.now(UTC)
        .strftime(
            "%Y-%m"
        )
    )




# ============================================================
# COMPLETE REPORT DATA
# ============================================================

def build_report_data(
    project_code,
    requested_period=None,
):

    with get_connection() as connection:

        reporting_period = (
            resolve_reporting_period(
                connection,
                requested_period,
                project_code,
            )
        )


        scorecard = (
            get_project_scorecard(
                connection,
                project_code,
                reporting_period,
            )
        )


        change_control = (
            get_change_control(
                connection,
                project_code,
                scorecard,
            )
        )


        risk_control = (
            get_risk_control(
                connection,
                project_code,
            )
        )


        cashflow = (
            get_cashflow(
                connection,
                project_code,
            )
        )


        commentary = (
            get_workflow_commentary(
                connection,
                project_code,
                reporting_period,
            )
        )


        actions = (
            get_workflow_actions(
                connection,
                project_code,
            )
        )


        schedule = (
            get_schedule_summary(
                connection,
                project_code,
            )
        )


        procurement = (
            get_procurement_summary(
                connection,
                project_code,
            )
        )


    project_name = (
        scorecard.get(
            "project_name"
        )
        or
        project_code
    )


    report = {

        "generated_at":
            utc_timestamp(),

        "reporting_period":
            reporting_period,

        "project": {

            "project_code":
                project_code,

            "project_name":
                project_name,
        },

        "performance": {

            "actual_cost_to_date":
                safe_float(
                    scorecard.get(
                        "actual_cost_to_date"
                    )
                ),

            "committed_cost_to_date":
                safe_float(
                    scorecard.get(
                        "committed_cost_to_date"
                    )
                ),

            "forecast_at_completion":
                safe_float(
                    scorecard.get(
                        "forecast_at_completion"
                    )
                    or
                    scorecard.get(
                        "fac"
                    )
                ),

            "forecast_variance":
                safe_float(
                    scorecard.get(
                        "forecast_variance"
                    )
                ),

            "forecast_variance_pct":
                safe_float(
                    scorecard.get(
                        "forecast_variance_pct"
                    )
                ),

            "planned_progress_pct":
                safe_float(
                    scorecard.get(
                        "planned_progress_pct"
                    )
                ),

            "actual_progress_pct":
                safe_float(
                    scorecard.get(
                        "actual_progress_pct"
                    )
                ),

            "progress_variance_pct":
                safe_float(
                    scorecard.get(
                        "progress_variance_pct"
                    )
                ),

            "management_status":
                scorecard.get(
                    "management_status"
                ),

            "forecast_credibility_status":
                scorecard.get(
                    "forecast_credibility_status"
                ),

            "progress_implied_eac":
                safe_float(
                    scorecard.get(
                        "progress_implied_eac"
                    )
                ),

            "forecast_credibility_gap":
                safe_float(
                    scorecard.get(
                        "forecast_credibility_gap"
                    )
                ),

            "forecast_credibility_gap_pct":
                safe_float(
                    scorecard.get(
                        "forecast_credibility_gap_pct"
                    )
                ),
        },

        "change_control":
            change_control,

        "risk_control":
            risk_control,

        "cashflow":
            cashflow,

        "schedule":
            schedule,

        "procurement":
            procurement,

        "commentary":
            commentary,

        "actions":
            actions,

        "methodology_notes": [

            (
                "Forecast concern uses a "
                "progress-implied completion-cost "
                "proxy as a management cross-check. "
                "It is not formal earned-value EAC."
            ),

            (
                "Cash-flow planning is maintained "
                "separately from accounting cost "
                "and forecast-at-completion."
            ),

            (
                "Approved changes alter Current Budget. "
                "Pending changes remain exposure only."
            ),

            (
                "Proposed risks do not consume "
                "contingency until approved."
            ),

            (
                "The Management Attention Index "
                "is a prioritisation tool and not "
                "a probability of project failure."
            ),
        ],
    }


    return report


# ============================================================
# PDF HELPERS
# ============================================================

def pdf_footer(
    canvas,
    doc,
):

    canvas.saveState()

    canvas.setFont(
        "Helvetica",
        7,
    )

    canvas.setFillColor(
        colors.HexColor(
            "#64748B"
        )
    )

    canvas.drawString(
        15 * mm,
        8 * mm,
        "ControlSight | Project Controls Management Report",
    )

    canvas.drawRightString(
        A4[0] - 15 * mm,
        8 * mm,
        f"Page {doc.page}",
    )

    canvas.restoreState()


def create_pdf(
    report,
    output_path,
):

    if not REPORTLAB_AVAILABLE:

        raise HTTPException(
            status_code=500,
            detail=(
                "PDF export requires reportlab. "
                "Run: pip install reportlab"
            ),
        )


    doc = SimpleDocTemplate(
        str(
            output_path
        ),
        pagesize=A4,
        rightMargin=14 * mm,
        leftMargin=14 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
    )


    styles = getSampleStyleSheet()


    title_style = ParagraphStyle(
        "ControlSightTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=colors.HexColor(
            "#0F172A"
        ),
        spaceAfter=5 * mm,
    )


    subtitle_style = ParagraphStyle(
        "ControlSightSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        textColor=colors.HexColor(
            "#64748B"
        ),
        spaceAfter=5 * mm,
    )


    section_style = ParagraphStyle(
        "ControlSightSection",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=15,
        textColor=colors.HexColor(
            "#1E3A5F"
        ),
        spaceBefore=5 * mm,
        spaceAfter=3 * mm,
    )


    body_style = ParagraphStyle(
        "ControlSightBody",
        parent=styles["BodyText"],
        fontSize=8,
        leading=11,
        textColor=colors.HexColor(
            "#334155"
        ),
    )


    small_style = ParagraphStyle(
        "ControlSightSmall",
        parent=styles["BodyText"],
        fontSize=7,
        leading=9,
        textColor=colors.HexColor(
            "#475569"
        ),
    )


    story = []


    project = report[
        "project"
    ]


    story.append(
        Paragraph(
            "ControlSight Monthly Management Report",
            title_style,
        )
    )


    story.append(
        Paragraph(
            (
                f"<b>{project['project_code']}</b> · "
                f"{project['project_name']}<br/>"
                f"Reporting period: "
                f"{report['reporting_period']}<br/>"
                f"Generated: "
                f"{report['generated_at']}"
            ),
            subtitle_style,
        )
    )


    # --------------------------------------------------------
    # EXECUTIVE SCORECARD
    # --------------------------------------------------------

    story.append(
        Paragraph(
            "1. Executive Project Scorecard",
            section_style,
        )
    )


    performance = report[
        "performance"
    ]

    change = report[
        "change_control"
    ]


    scorecard_data = [

        [
            "Current Budget",
            money(
                change[
                    "current_budget"
                ]
            ),
            "Actual Cost",
            money(
                performance[
                    "actual_cost_to_date"
                ]
            ),
        ],

        [
            "Committed Cost",
            money(
                performance[
                    "committed_cost_to_date"
                ]
            ),
            "FAC",
            money(
                performance[
                    "forecast_at_completion"
                ]
            ),
        ],

        [
            "Forecast Variance",
            money(
                performance[
                    "forecast_variance"
                ]
            ),
            "Forecast Variance %",
            percentage(
                performance[
                    "forecast_variance_pct"
                ]
            ),
        ],

        [
            "Planned Progress",
            percentage(
                performance[
                    "planned_progress_pct"
                ]
            ),
            "Actual Progress",
            percentage(
                performance[
                    "actual_progress_pct"
                ]
            ),
        ],

        [
            "Progress Variance",
            percentage(
                performance[
                    "progress_variance_pct"
                ]
            ),
            "Management Status",
            text_value(
                performance[
                    "management_status"
                ]
            ),
        ],
    ]


    scorecard_table = Table(
        scorecard_data,
        colWidths=[
            35 * mm,
            45 * mm,
            35 * mm,
            45 * mm,
        ],
    )


    scorecard_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, -1),
                    colors.HexColor(
                        "#F8FAFC"
                    ),
                ),
                (
                    "BOX",
                    (0, 0),
                    (-1, -1),
                    0.4,
                    colors.HexColor(
                        "#CBD5E1"
                    ),
                ),
                (
                    "INNERGRID",
                    (0, 0),
                    (-1, -1),
                    0.3,
                    colors.HexColor(
                        "#E2E8F0"
                    ),
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, -1),
                    "Helvetica",
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (0, -1),
                    "Helvetica-Bold",
                ),
                (
                    "FONTNAME",
                    (2, 0),
                    (2, -1),
                    "Helvetica-Bold",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    8,
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
            ]
        )
    )


    story.append(
        scorecard_table
    )


    # --------------------------------------------------------
    # FORECAST CREDIBILITY
    # --------------------------------------------------------

    story.append(
        Paragraph(
            "2. Forecast Concern",
            section_style,
        )
    )


    forecast_table = Table(
        [
            [
                "Reported FAC",
                money(
                    performance[
                        "forecast_at_completion"
                    ]
                ),
            ],
            [
                "Progress-Implied EAC",
                money(
                    performance[
                        "progress_implied_eac"
                    ]
                ),
            ],
            [
                "Forecast Gap",
                money(
                    performance[
                        "forecast_credibility_gap"
                    ]
                ),
            ],
            [
                "Forecast Gap %",
                percentage(
                    performance[
                        "forecast_credibility_gap_pct"
                    ]
                ),
            ],
            [
                "Concern Status",
                text_value(
                    performance[
                        "forecast_credibility_status"
                    ]
                ),
            ],
        ],
        colWidths=[
            65 * mm,
            95 * mm,
        ],
    )


    forecast_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (0, -1),
                    colors.HexColor(
                        "#EFF6FF"
                    ),
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (0, -1),
                    "Helvetica-Bold",
                ),
                (
                    "BOX",
                    (0, 0),
                    (-1, -1),
                    0.4,
                    colors.HexColor(
                        "#BFDBFE"
                    ),
                ),
                (
                    "INNERGRID",
                    (0, 0),
                    (-1, -1),
                    0.3,
                    colors.HexColor(
                        "#DBEAFE"
                    ),
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    8,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
            ]
        )
    )


    story.append(
        forecast_table
    )


    # --------------------------------------------------------
    # CHANGE CONTROL
    # --------------------------------------------------------

    story.append(
        Paragraph(
            "3. Budget & Change Control",
            section_style,
        )
    )


    budget_data = [

        [
            "Original Budget",
            money(
                change[
                    "original_budget"
                ]
            ),
        ],

        [
            "Approved Changes",
            money(
                change[
                    "approved_changes"
                ]
            ),
        ],

        [
            "Current Budget",
            money(
                change[
                    "current_budget"
                ]
            ),
        ],

        [
            "Pending Changes",
            money(
                change[
                    "pending_changes"
                ]
            ),
        ],

        [
            "Potential Budget",
            money(
                change[
                    "potential_budget"
                ]
            ),
        ],
    ]


    budget_table = Table(
        budget_data,
        colWidths=[
            65 * mm,
            95 * mm,
        ],
    )


    budget_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (0, -1),
                    colors.HexColor(
                        "#F8FAFC"
                    ),
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (0, -1),
                    "Helvetica-Bold",
                ),
                (
                    "BOX",
                    (0, 0),
                    (-1, -1),
                    0.4,
                    colors.HexColor(
                        "#CBD5E1"
                    ),
                ),
                (
                    "INNERGRID",
                    (0, 0),
                    (-1, -1),
                    0.3,
                    colors.HexColor(
                        "#E2E8F0"
                    ),
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    8,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
            ]
        )
    )


    story.append(
        budget_table
    )


    changes = change[
        "changes"
    ]


    if changes:

        story.append(
            Spacer(
                1,
                3 * mm,
            )
        )


        change_rows = [

            [
                "Reference",
                "Category",
                "Amount",
                "Status",
                "Title",
            ]
        ]


        for row in changes:

            change_rows.append(
                [
                    text_value(
                        row.get(
                            "change_reference"
                        )
                    ),
                    text_value(
                        row.get(
                            "category"
                        )
                    ),
                    money(
                        row.get(
                            "change_amount"
                        )
                    ),
                    text_value(
                        row.get(
                            "status"
                        )
                    ),
                    Paragraph(
                        text_value(
                            row.get(
                                "title"
                            )
                        ),
                        small_style,
                    ),
                ]
            )


        change_table = Table(
            change_rows,
            repeatRows=1,
            colWidths=[
                28 * mm,
                30 * mm,
                26 * mm,
                24 * mm,
                52 * mm,
            ],
        )


        change_table.setStyle(
            standard_pdf_table_style()
        )


        story.append(
            change_table
        )


    # --------------------------------------------------------
    # RISK / CONTINGENCY
    # --------------------------------------------------------

    story.append(
        Paragraph(
            "4. Contingency & Risk Drawdown",
            section_style,
        )
    )


    risk = report[
        "risk_control"
    ]


    risk_summary = Table(
        [
            [
                "Contingency",
                money(
                    risk[
                        "contingency_budget"
                    ]
                ),
                "Proposed Exposure",
                money(
                    risk[
                        "proposed_exposure"
                    ]
                ),
            ],
            [
                "Approved not Drawn",
                money(
                    risk[
                        "approved_not_drawn"
                    ]
                ),
                "Drawn",
                money(
                    risk[
                        "drawn"
                    ]
                ),
            ],
            [
                "Remaining",
                money(
                    risk[
                        "remaining"
                    ]
                ),
                "Utilisation",
                percentage(
                    risk[
                        "utilisation_pct"
                    ]
                ),
            ],
        ],
        colWidths=[
            35 * mm,
            45 * mm,
            35 * mm,
            45 * mm,
        ],
    )


    risk_summary.setStyle(
        standard_pdf_table_style(
            header=False
        )
    )


    story.append(
        risk_summary
    )


    risks = risk[
        "risks"
    ]


    if risks:

        story.append(
            Spacer(
                1,
                3 * mm,
            )
        )


        risk_rows = [

            [
                "Reference",
                "Category",
                "Amount",
                "Status",
                "Owner",
                "Risk",
            ]
        ]


        for row in risks:

            risk_rows.append(
                [
                    text_value(
                        row.get(
                            "risk_reference"
                        )
                    ),
                    text_value(
                        row.get(
                            "category"
                        )
                    ),
                    money(
                        row.get(
                            "amount"
                        )
                    ),
                    text_value(
                        row.get(
                            "status"
                        )
                    ),
                    text_value(
                        row.get(
                            "owner"
                        )
                    ),
                    Paragraph(
                        text_value(
                            row.get(
                                "title"
                            )
                        ),
                        small_style,
                    ),
                ]
            )


        risk_table = Table(
            risk_rows,
            repeatRows=1,
            colWidths=[
                27 * mm,
                24 * mm,
                23 * mm,
                21 * mm,
                27 * mm,
                38 * mm,
            ],
        )


        risk_table.setStyle(
            standard_pdf_table_style()
        )


        story.append(
            risk_table
        )


    # --------------------------------------------------------
    # SCHEDULE / PROCUREMENT
    # --------------------------------------------------------

    story.append(
        Paragraph(
            "5. Schedule & Procurement",
            section_style,
        )
    )


    schedule = report[
        "schedule"
    ]

    procurement = report[
        "procurement"
    ]


    sp_table = Table(
        [
            [
                "Schedule Metric",
                "Value",
                "Procurement Metric",
                "Value",
            ],
            [
                "Total Activities",
                text_value(
                    schedule.get(
                        "total_activities"
                    )
                ),
                "Total Items",
                text_value(
                    procurement.get(
                        "total_items"
                    )
                ),
            ],
            [
                "Delayed Activities",
                text_value(
                    schedule.get(
                        "delayed_activities"
                    )
                ),
                "Delayed Items",
                text_value(
                    procurement.get(
                        "delayed_items"
                    )
                ),
            ],
            [
                "Critical Activities",
                text_value(
                    schedule.get(
                        "critical_activities"
                    )
                ),
                "Critical Delays",
                text_value(
                    procurement.get(
                        "critical_delays"
                    )
                ),
            ],
            [
                "Critical Delayed",
                text_value(
                    schedule.get(
                        "critical_delayed"
                    )
                ),
                "",
                "",
            ],
        ],
        colWidths=[
            45 * mm,
            25 * mm,
            45 * mm,
            25 * mm,
        ],
    )


    sp_table.setStyle(
        standard_pdf_table_style()
    )


    story.append(
        sp_table
    )


    # --------------------------------------------------------
    # CASH FLOW
    # --------------------------------------------------------

    story.append(
        Paragraph(
            "6. Cash-Flow Position",
            section_style,
        )
    )


    cash = report[
        "cashflow"
    ]


    cash_summary = Table(
        [
            [
                "Forecast Cash",
                money(
                    cash[
                        "total_forecast"
                    ]
                ),
                "Actual Cash",
                money(
                    cash[
                        "total_actual"
                    ]
                ),
            ],
            [
                "Committed Cash",
                money(
                    cash[
                        "total_committed"
                    ]
                ),
                "Cash Variance",
                money(
                    cash[
                        "total_variance"
                    ]
                ),
            ],
        ],
        colWidths=[
            35 * mm,
            45 * mm,
            35 * mm,
            45 * mm,
        ],
    )


    cash_summary.setStyle(
        standard_pdf_table_style(
            header=False
        )
    )


    story.append(
        cash_summary
    )


    if cash[
        "periods"
    ]:

        story.append(
            Spacer(
                1,
                3 * mm,
            )
        )


        cash_rows = [

            [
                "Period",
                "Forecast",
                "Actual",
                "Committed",
                "Variance",
                "Variance %",
            ]
        ]


        for row in cash[
            "periods"
        ]:

            cash_rows.append(
                [
                    text_value(
                        row.get(
                            "reporting_period"
                        )
                    ),
                    money(
                        row.get(
                            "forecast_cash"
                        )
                    ),
                    money(
                        row.get(
                            "actual_cash"
                        )
                    ),
                    money(
                        row.get(
                            "committed_cash"
                        )
                    ),
                    money(
                        row.get(
                            "variance"
                        )
                    ),
                    percentage(
                        row.get(
                            "variance_pct"
                        )
                    ),
                ]
            )


        cash_table = Table(
            cash_rows,
            repeatRows=1,
            colWidths=[
                25 * mm,
                28 * mm,
                28 * mm,
                28 * mm,
                28 * mm,
                23 * mm,
            ],
        )


        cash_table.setStyle(
            standard_pdf_table_style()
        )


        story.append(
            cash_table
        )


    # --------------------------------------------------------
    # COMMENTARY
    # --------------------------------------------------------

    story.append(
        Paragraph(
            "7. Management Commentary",
            section_style,
        )
    )


    commentary = report[
        "commentary"
    ]


    commentary_items = [

        (
            "Cost",
            commentary.get(
                "cost_commentary"
            )
        ),

        (
            "Progress",
            commentary.get(
                "progress_commentary"
            )
        ),

        (
            "Schedule",
            commentary.get(
                "schedule_commentary"
            )
        ),

        (
            "Procurement",
            commentary.get(
                "procurement_commentary"
            )
        ),

        (
            "Forecast Concern",
            commentary.get(
                "forecast_credibility_note"
            )
        ),
    ]


    for label, value in commentary_items:

        story.append(
            Paragraph(
                (
                    f"<b>{label}:</b> "
                    f"{text_value(value)}"
                ),
                body_style,
            )
        )

        story.append(
            Spacer(
                1,
                2 * mm,
            )
        )


    # --------------------------------------------------------
    # ACTION LOG
    # --------------------------------------------------------

    story.append(
        Paragraph(
            "8. Management Action Log",
            section_style,
        )
    )


    actions = report[
        "actions"
    ]


    if actions:

        action_rows = [

            [
                "Category",
                "Action",
                "Owner",
                "Due",
                "Priority",
                "Status",
            ]
        ]


        for row in actions:

            action_text = (
                row.get(
                    "action"
                )
                or
                row.get(
                    "action_text"
                )
                or
                row.get(
                    "description"
                )
                or
                row.get(
                    "issue"
                )
            )


            owner = (
                row.get(
                    "owner"
                )
                or
                row.get(
                    "owner_role"
                )
            )


            due_date = (
                row.get(
                    "due_date"
                )
                or
                row.get(
                    "due"
                )
            )


            action_rows.append(
                [
                    text_value(
                        row.get(
                            "category"
                        )
                    ),
                    Paragraph(
                        text_value(
                            action_text
                        ),
                        small_style,
                    ),
                    text_value(
                        owner
                    ),
                    text_value(
                        due_date
                    ),
                    text_value(
                        row.get(
                            "priority"
                        )
                    ),
                    text_value(
                        row.get(
                            "status"
                        )
                    ),
                ]
            )


        action_table = Table(
            action_rows,
            repeatRows=1,
            colWidths=[
                25 * mm,
                55 * mm,
                30 * mm,
                22 * mm,
                20 * mm,
                23 * mm,
            ],
        )


        action_table.setStyle(
            standard_pdf_table_style()
        )


        story.append(
            action_table
        )

    else:

        story.append(
            Paragraph(
                "No management actions recorded.",
                body_style,
            )
        )


    # --------------------------------------------------------
    # METHODOLOGY
    # --------------------------------------------------------

    story.append(
        Paragraph(
            "9. Methodology & Control Notes",
            section_style,
        )
    )


    for note in report[
        "methodology_notes"
    ]:

        story.append(
            Paragraph(
                f"• {note}",
                body_style,
            )
        )

        story.append(
            Spacer(
                1,
                1.5 * mm,
            )
        )


    doc.build(
        story,
        onFirstPage=pdf_footer,
        onLaterPages=pdf_footer,
    )


def standard_pdf_table_style(
    header=True,
):

    commands = [

        (
            "BOX",
            (0, 0),
            (-1, -1),
            0.4,
            colors.HexColor(
                "#CBD5E1"
            ),
        ),

        (
            "INNERGRID",
            (0, 0),
            (-1, -1),
            0.3,
            colors.HexColor(
                "#E2E8F0"
            ),
        ),

        (
            "FONTSIZE",
            (0, 0),
            (-1, -1),
            7,
        ),

        (
            "VALIGN",
            (0, 0),
            (-1, -1),
            "TOP",
        ),

        (
            "TOPPADDING",
            (0, 0),
            (-1, -1),
            5,
        ),

        (
            "BOTTOMPADDING",
            (0, 0),
            (-1, -1),
            5,
        ),
    ]


    if header:

        commands.extend(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor(
                        "#1E3A5F"
                    ),
                ),

                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.white,
                ),

                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
            ]
        )


    return TableStyle(
        commands
    )


# ============================================================
# EXCEL HELPERS
# ============================================================

def style_excel_sheet(
    worksheet,
    title,
):

    worksheet.sheet_view.showGridLines = (
        False
    )


    worksheet["A1"] = title


    worksheet["A1"].font = Font(
        size=18,
        bold=True,
    )


    worksheet["A1"].fill = PatternFill(
        fill_type="solid",
        fgColor="1E3A5F",
    )


    worksheet["A1"].font = Font(
        size=18,
        bold=True,
        color="FFFFFF",
    )


    worksheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=6,
    )


    worksheet.row_dimensions[
        1
    ].height = 28


def excel_header_style(
    cell,
):

    cell.font = Font(
        bold=True,
        color="FFFFFF",
    )

    cell.fill = PatternFill(
        fill_type="solid",
        fgColor="1E3A5F",
    )

    cell.alignment = Alignment(
        vertical="center",
    )


def excel_section_style(
    cell,
):

    cell.font = Font(
        bold=True,
        color="1E3A5F",
        size=12,
    )


def autosize_excel(
    worksheet,
    max_width=45,
):

    for column_cells in (
        worksheet.columns
    ):

        length = 0

        column_letter = (
            get_column_letter(
                column_cells[
                    0
                ].column
            )
        )


        for cell in column_cells:

            if cell.value is None:
                continue

            value_length = len(
                str(
                    cell.value
                )
            )

            length = max(
                length,
                value_length,
            )


        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(
                length + 2,
                10,
            ),
            max_width,
        )


def create_excel(
    report,
    output_path,
):

    if not OPENPYXL_AVAILABLE:

        raise HTTPException(
            status_code=500,
            detail=(
                "Excel export requires openpyxl. "
                "Run: pip install openpyxl"
            ),
        )


    workbook = Workbook()


    default_sheet = (
        workbook.active
    )

    workbook.remove(
        default_sheet
    )


    # ========================================================
    # EXECUTIVE SUMMARY
    # ========================================================

    ws = workbook.create_sheet(
        "Executive Summary"
    )


    style_excel_sheet(
        ws,
        "ControlSight Monthly Management Report",
    )


    project = report[
        "project"
    ]

    performance = report[
        "performance"
    ]

    change = report[
        "change_control"
    ]


    ws["A3"] = "Project"
    ws["B3"] = (
        f"{project['project_code']} · "
        f"{project['project_name']}"
    )

    ws["A4"] = "Reporting Period"
    ws["B4"] = report[
        "reporting_period"
    ]

    ws["A5"] = "Generated"
    ws["B5"] = report[
        "generated_at"
    ]


    for cell in [
        "A3",
        "A4",
        "A5",
    ]:

        ws[cell].font = Font(
            bold=True
        )


    executive_rows = [

        (
            "Current Budget",
            change[
                "current_budget"
            ],
        ),

        (
            "Actual Cost",
            performance[
                "actual_cost_to_date"
            ],
        ),

        (
            "Committed Cost",
            performance[
                "committed_cost_to_date"
            ],
        ),

        (
            "FAC",
            performance[
                "forecast_at_completion"
            ],
        ),

        (
            "Forecast Variance",
            performance[
                "forecast_variance"
            ],
        ),

        (
            "Forecast Variance %",
            performance[
                "forecast_variance_pct"
            ],
        ),

        (
            "Planned Progress %",
            performance[
                "planned_progress_pct"
            ],
        ),

        (
            "Actual Progress %",
            performance[
                "actual_progress_pct"
            ],
        ),

        (
            "Progress Variance",
            performance[
                "progress_variance_pct"
            ],
        ),

        (
            "Forecast Concern",
            performance[
                "forecast_credibility_status"
            ],
        ),

        (
            "Progress-Implied EAC",
            performance[
                "progress_implied_eac"
            ],
        ),

        (
            "Forecast Concern Gap",
            performance[
                "forecast_credibility_gap"
            ],
        ),
    ]


    row_number = 8


    for label, value in executive_rows:

        ws.cell(
            row=row_number,
            column=1,
            value=label,
        ).font = Font(
            bold=True
        )

        ws.cell(
            row=row_number,
            column=2,
            value=value,
        )

        row_number += 1


    autosize_excel(
        ws
    )


    # ========================================================
    # CHANGE CONTROL
    # ========================================================

    ws = workbook.create_sheet(
        "Change Control"
    )


    style_excel_sheet(
        ws,
        "Budget & Change Control",
    )


    summary_rows = [

        (
            "Original Budget",
            change[
                "original_budget"
            ],
        ),

        (
            "Approved Changes",
            change[
                "approved_changes"
            ],
        ),

        (
            "Current Budget",
            change[
                "current_budget"
            ],
        ),

        (
            "Pending Changes",
            change[
                "pending_changes"
            ],
        ),

        (
            "Potential Budget",
            change[
                "potential_budget"
            ],
        ),
    ]


    row_number = 3


    for label, value in summary_rows:

        ws.cell(
            row=row_number,
            column=1,
            value=label,
        ).font = Font(
            bold=True
        )

        ws.cell(
            row=row_number,
            column=2,
            value=value,
        )

        row_number += 1


    row_number += 2


    headers = [
        "Reference",
        "Category",
        "Title",
        "Description",
        "Amount",
        "Status",
        "Approval Date",
        "Requested By",
        "Approved By",
    ]


    for column, header in enumerate(
        headers,
        start=1,
    ):

        cell = ws.cell(
            row=row_number,
            column=column,
            value=header,
        )

        excel_header_style(
            cell
        )


    row_number += 1


    for change_row in change[
        "changes"
    ]:

        values = [

            change_row.get(
                "change_reference"
            ),

            change_row.get(
                "category"
            ),

            change_row.get(
                "title"
            ),

            change_row.get(
                "description"
            ),

            safe_float(
                change_row.get(
                    "change_amount"
                )
            ),

            change_row.get(
                "status"
            ),

            change_row.get(
                "approval_date"
            ),

            change_row.get(
                "requested_by"
            ),

            change_row.get(
                "approved_by"
            ),
        ]


        for column, value in enumerate(
            values,
            start=1,
        ):

            ws.cell(
                row=row_number,
                column=column,
                value=value,
            )


        row_number += 1


    autosize_excel(
        ws
    )


    # ========================================================
    # RISK & CONTINGENCY
    # ========================================================

    risk = report[
        "risk_control"
    ]


    ws = workbook.create_sheet(
        "Risk & Contingency"
    )


    style_excel_sheet(
        ws,
        "Contingency & Risk Drawdown",
    )


    risk_summary_rows = [

        (
            "Contingency Budget",
            risk[
                "contingency_budget"
            ],
        ),

        (
            "Proposed Exposure",
            risk[
                "proposed_exposure"
            ],
        ),

        (
            "Approved not Drawn",
            risk[
                "approved_not_drawn"
            ],
        ),

        (
            "Drawn",
            risk[
                "drawn"
            ],
        ),

        (
            "Committed",
            risk[
                "committed"
            ],
        ),

        (
            "Remaining",
            risk[
                "remaining"
            ],
        ),

        (
            "Utilisation %",
            risk[
                "utilisation_pct"
            ],
        ),
    ]


    row_number = 3


    for label, value in risk_summary_rows:

        ws.cell(
            row=row_number,
            column=1,
            value=label,
        ).font = Font(
            bold=True
        )

        ws.cell(
            row=row_number,
            column=2,
            value=value,
        )

        row_number += 1


    row_number += 2


    headers = [
        "Reference",
        "Category",
        "Title",
        "Description",
        "Amount",
        "Status",
        "Owner",
        "Event Date",
        "Approved By",
    ]


    for column, header in enumerate(
        headers,
        start=1,
    ):

        cell = ws.cell(
            row=row_number,
            column=column,
            value=header,
        )

        excel_header_style(
            cell
        )


    row_number += 1


    for risk_row in risk[
        "risks"
    ]:

        values = [

            risk_row.get(
                "risk_reference"
            ),

            risk_row.get(
                "category"
            ),

            risk_row.get(
                "title"
            ),

            risk_row.get(
                "description"
            ),

            safe_float(
                risk_row.get(
                    "amount"
                )
            ),

            risk_row.get(
                "status"
            ),

            risk_row.get(
                "owner"
            ),

            risk_row.get(
                "event_date"
            ),

            risk_row.get(
                "approved_by"
            ),
        ]


        for column, value in enumerate(
            values,
            start=1,
        ):

            ws.cell(
                row=row_number,
                column=column,
                value=value,
            )


        row_number += 1


    autosize_excel(
        ws
    )


    # ========================================================
    # CASH FLOW
    # ========================================================

    cash = report[
        "cashflow"
    ]


    ws = workbook.create_sheet(
        "Cash Flow"
    )


    style_excel_sheet(
        ws,
        "Monthly Cash-Flow Position",
    )


    headers = [
        "Reporting Period",
        "Forecast Cash",
        "Actual Cash",
        "Committed Cash",
        "Variance",
        "Variance %",
        "Cumulative Forecast",
        "Cumulative Actual",
        "Notes",
    ]


    row_number = 3


    for column, header in enumerate(
        headers,
        start=1,
    ):

        cell = ws.cell(
            row=row_number,
            column=column,
            value=header,
        )

        excel_header_style(
            cell
        )


    row_number += 1


    for cash_row in cash[
        "periods"
    ]:

        values = [

            cash_row.get(
                "reporting_period"
            ),

            safe_float(
                cash_row.get(
                    "forecast_cash"
                )
            ),

            safe_float(
                cash_row.get(
                    "actual_cash"
                )
            ),

            safe_float(
                cash_row.get(
                    "committed_cash"
                )
            ),

            safe_float(
                cash_row.get(
                    "variance"
                )
            ),

            safe_float(
                cash_row.get(
                    "variance_pct"
                )
            ),

            safe_float(
                cash_row.get(
                    "cumulative_forecast"
                )
            ),

            safe_float(
                cash_row.get(
                    "cumulative_actual"
                )
            ),

            cash_row.get(
                "notes"
            ),
        ]


        for column, value in enumerate(
            values,
            start=1,
        ):

            ws.cell(
                row=row_number,
                column=column,
                value=value,
            )


        row_number += 1


    autosize_excel(
        ws
    )


    # ========================================================
    # SCHEDULE / PROCUREMENT
    # ========================================================

    ws = workbook.create_sheet(
        "Schedule Procurement"
    )


    style_excel_sheet(
        ws,
        "Schedule & Procurement Summary",
    )


    schedule = report[
        "schedule"
    ]

    procurement = report[
        "procurement"
    ]


    values = [

        (
            "Total Activities",
            schedule.get(
                "total_activities"
            ),
        ),

        (
            "Delayed Activities",
            schedule.get(
                "delayed_activities"
            ),
        ),

        (
            "Critical Activities",
            schedule.get(
                "critical_activities"
            ),
        ),

        (
            "Critical Delayed",
            schedule.get(
                "critical_delayed"
            ),
        ),

        (
            "Procurement Items",
            procurement.get(
                "total_items"
            ),
        ),

        (
            "Delayed Procurement",
            procurement.get(
                "delayed_items"
            ),
        ),

        (
            "Critical Procurement Delays",
            procurement.get(
                "critical_delays"
            ),
        ),
    ]


    row_number = 3


    for label, value in values:

        ws.cell(
            row=row_number,
            column=1,
            value=label,
        ).font = Font(
            bold=True
        )

        ws.cell(
            row=row_number,
            column=2,
            value=value,
        )

        row_number += 1


    autosize_excel(
        ws
    )


    # ========================================================
    # COMMENTARY
    # ========================================================

    ws = workbook.create_sheet(
        "Commentary"
    )


    style_excel_sheet(
        ws,
        "Management Commentary",
    )


    commentary = report[
        "commentary"
    ]


    commentary_rows = [

        (
            "Cost",
            commentary.get(
                "cost_commentary"
            ),
        ),

        (
            "Progress",
            commentary.get(
                "progress_commentary"
            ),
        ),

        (
            "Schedule",
            commentary.get(
                "schedule_commentary"
            ),
        ),

        (
            "Procurement",
            commentary.get(
                "procurement_commentary"
            ),
        ),

        (
            "Forecast Concern",
            commentary.get(
                "forecast_credibility_note"
            ),
        ),
    ]


    row_number = 3


    for label, value in commentary_rows:

        ws.cell(
            row=row_number,
            column=1,
            value=label,
        ).font = Font(
            bold=True
        )

        ws.cell(
            row=row_number,
            column=2,
            value=text_value(
                value
            ),
        )

        ws.cell(
            row=row_number,
            column=2,
        ).alignment = Alignment(
            wrap_text=True,
            vertical="top",
        )

        row_number += 1


    autosize_excel(
        ws
    )


    # ========================================================
    # ACTION LOG
    # ========================================================

    ws = workbook.create_sheet(
        "Action Log"
    )


    style_excel_sheet(
        ws,
        "Management Action Log",
    )


    headers = [
        "Category",
        "Issue / Action",
        "Owner",
        "Due Date",
        "Priority",
        "Status",
    ]


    row_number = 3


    for column, header in enumerate(
        headers,
        start=1,
    ):

        cell = ws.cell(
            row=row_number,
            column=column,
            value=header,
        )

        excel_header_style(
            cell
        )


    row_number += 1


    for action in report[
        "actions"
    ]:

        action_text = (
            action.get(
                "action"
            )
            or
            action.get(
                "action_text"
            )
            or
            action.get(
                "description"
            )
            or
            action.get(
                "issue"
            )
        )


        owner = (
            action.get(
                "owner"
            )
            or
            action.get(
                "owner_role"
            )
        )


        due_date = (
            action.get(
                "due_date"
            )
            or
            action.get(
                "due"
            )
        )


        values = [

            action.get(
                "category"
            ),

            action_text,

            owner,

            due_date,

            action.get(
                "priority"
            ),

            action.get(
                "status"
            ),
        ]


        for column, value in enumerate(
            values,
            start=1,
        ):

            ws.cell(
                row=row_number,
                column=column,
                value=value,
            )


        row_number += 1


    autosize_excel(
        ws
    )


    # ========================================================
    # METHODOLOGY
    # ========================================================

    ws = workbook.create_sheet(
        "Methodology"
    )


    style_excel_sheet(
        ws,
        "Control Methodology Notes",
    )


    row_number = 3


    for index, note in enumerate(
        report[
            "methodology_notes"
        ],
        start=1,
    ):

        ws.cell(
            row=row_number,
            column=1,
            value=index,
        ).font = Font(
            bold=True
        )

        ws.cell(
            row=row_number,
            column=2,
            value=note,
        ).alignment = Alignment(
            wrap_text=True,
            vertical="top",
        )

        row_number += 1


    autosize_excel(
        ws,
        max_width=80,
    )


    # ========================================================
    # NUMBER FORMATTING
    # ========================================================

    for sheet in (
        workbook.worksheets
    ):

        for row in sheet.iter_rows():

            for cell in row:

                if (
                    isinstance(
                        cell.value,
                        float,
                    )
                    or
                    isinstance(
                        cell.value,
                        int,
                    )
                ):

                    cell.number_format = (
                        '#,##0.00'
                    )


    workbook.save(
        output_path
    )


# ============================================================
# API
# ============================================================

@router.get(
    "/health/status"
)
def report_pack_health():

    return {

        "status":
            "ok",

        "module":
            "Monthly Report Pack",

        "database_available":
            DATABASE_PATH.exists(),

        "pdf_export":
            REPORTLAB_AVAILABLE,

        "excel_export":
            OPENPYXL_AVAILABLE,

        "output_directory":
            str(
                REPORT_OUTPUT_DIR
            ),
    }


@router.get(
    "/preview/{project_code}"
)
def preview_report_pack(
    project_code: str,
    reporting_period: str | None = None,
):

    return build_report_data(
        project_code,
        reporting_period,
    )


@router.post(
    "/generate"
)
def generate_report_pack(
    request: ReportPackRequest,
):

    requested_formats = {
        item.lower().strip()
        for item in request.formats
    }


    invalid_formats = (
        requested_formats
        -
        SUPPORTED_FORMATS
    )


    if invalid_formats:

        raise HTTPException(
            status_code=422,
            detail=(
                "Unsupported formats: "
                f"{sorted(invalid_formats)}"
            ),
        )


    if not requested_formats:

        raise HTTPException(
            status_code=422,
            detail=(
                "At least one export format "
                "is required."
            ),
        )


    report = build_report_data(
        request.project_code,
        request.reporting_period,
    )


    project_code = (
        safe_filename_part(
            request.project_code
        )
    )


    reporting_period = (
        safe_filename_part(
            report[
                "reporting_period"
            ]
        )
    )


    timestamp = (
        datetime.now(UTC)
        .strftime(
            "%Y%m%d_%H%M%S"
        )
    )


    base_filename = (
        f"ControlSight_"
        f"{project_code}_"
        f"{reporting_period}_"
        f"{timestamp}"
    )


    generated_files = []


    if "pdf" in requested_formats:

        pdf_path = (
            REPORT_OUTPUT_DIR
            /
            f"{base_filename}.pdf"
        )


        create_pdf(
            report,
            pdf_path,
        )


        generated_files.append(
            {

                "format":
                    "pdf",

                "filename":
                    pdf_path.name,

                "download_url":
                    (
                        "/api/report-pack/files/"
                        f"{pdf_path.name}"
                    ),
            }
        )


    if "xlsx" in requested_formats:

        excel_path = (
            REPORT_OUTPUT_DIR
            /
            f"{base_filename}.xlsx"
        )


        create_excel(
            report,
            excel_path,
        )


        generated_files.append(
            {

                "format":
                    "xlsx",

                "filename":
                    excel_path.name,

                "download_url":
                    (
                        "/api/report-pack/files/"
                        f"{excel_path.name}"
                    ),
            }
        )


    return {

        "message":
            "Monthly report pack generated.",

        "project_code":
            request.project_code,

        "reporting_period":
            report[
                "reporting_period"
            ],

        "generated_at":
            report[
                "generated_at"
            ],

        "files":
            generated_files,
    }


@router.get(
    "/files/{filename}"
)
def download_report_pack_file(
    filename: str,
):

    safe_name = Path(
        filename
    ).name


    if safe_name != filename:

        raise HTTPException(
            status_code=400,
            detail=(
                "Invalid filename."
            ),
        )


    path = (
        REPORT_OUTPUT_DIR
        /
        safe_name
    )


    if not path.exists():

        raise HTTPException(
            status_code=404,
            detail=(
                "Report pack file not found."
            ),
        )


    suffix = (
        path.suffix.lower()
    )


    if suffix == ".pdf":

        media_type = (
            "application/pdf"
        )

    elif suffix == ".xlsx":

        media_type = (
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        )

    else:

        media_type = (
            "application/octet-stream"
        )


    return FileResponse(
        path=str(
            path
        ),
        filename=path.name,
        media_type=media_type,
    )


@router.get(
    "/files"
)
def list_report_pack_files():

    files = []


    for path in sorted(
        REPORT_OUTPUT_DIR.glob(
            "ControlSight_*"
        ),
        key=lambda item:
            item.stat().st_mtime,
        reverse=True,
    ):

        files.append(
            {

                "filename":
                    path.name,

                "size_bytes":
                    path.stat().st_size,

                "modified_at":
                    datetime.fromtimestamp(
                        path.stat().st_mtime
                    ).isoformat(),

                "download_url":
                    (
                        "/api/report-pack/files/"
                        f"{path.name}"
                    ),
            }
        )


    return {

        "count":
            len(
                files
            ),

        "files":
            files,
    }
